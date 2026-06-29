import re
from flask_socketio import SocketIO
from flask import Flask, render_template, request, Response, jsonify, session, redirect, url_for
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from flask_cors import CORS
import paho.mqtt.client as mqtt
import sqlite3
import json
import requests
from datetime import datetime
from werkzeug.security import generate_password_hash, check_password_hash

app = Flask(__name__)
app.secret_key = "evoluzn_secret_key_2024"
CORS(app)
socketio = SocketIO(app, cors_allowed_origins="*", async_mode="threading")

DB_PATH = "sensor_data.db"

BROKER = "evoluzn.org"
PORT   = 18889
#BROKER   = "broker.emqx.io"
#PORT     = 1883
USERNAME = "evzin_led"
PASSWORD = "63I9YhMaXpa49Eb"

# ── Global: stores the URL entered in Configure API modal ──
configured_url = {"url": ""}

def load_configured_url():
    """Load previously saved URL from DB into memory on startup."""
    try:
        conn = connect_db()
        cursor = conn.cursor()
        cursor.execute("SELECT url FROM configured_urls ORDER BY id DESC LIMIT 1")
        row = cursor.fetchone()
        cursor.close()
        conn.close()
        if row:
            configured_url["url"] = row["url"]
            print(f"✅ Loaded configured URL from DB: {configured_url['url']}")
    except Exception as e:
        print(f"⚠️ Could not load configured URL: {e}")


def connect_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def create_tables():
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS gateway_data (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            device_id TEXT,
            rs485_1 REAL,
            rs485_2 REAL,
            rs232_1 REAL,
            rs232_2 REAL,
            device_timestamp TEXT,
            timestamp DATETIME DEFAULT (datetime('now','localtime'))
        );
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            email TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            rtu_no TEXT,
            temperature INTEGER DEFAULT 0,
            voltage INTEGER DEFAULT 0,
            current INTEGER DEFAULT 0,
            power INTEGER DEFAULT 0
        );
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS configured_urls (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            company TEXT,
            url TEXT,
            updated_at DATETIME DEFAULT (datetime('now','localtime'))
        );
    """)
    conn.commit()
    cursor.close()
    conn.close()
    print("✅ Database tables are ready.")


def parse_gateway_payload(raw: str):
    try:
        raw = raw.strip()
        pattern = (
            r'device_id:\s*'
            r'([^\[\n:]+)\s*:?\s*'
            r'\[(.*?)\]\s*:?\s*'
            r'RS485_1:(.*?)\s*:?\s*'
            r'RS485_2:(.*?)\s*:?\s*'
            r'RS232_1:(.*?)\s*:?\s*'
            r'RS232_2:(.*)'
        )
        match = re.search(pattern, raw, re.DOTALL)
        if not match:
            print("❌ Payload format mismatch")
            return None

        def clean(val):
            val = val.strip().rstrip(':').rstrip('}')
            return None if val.lower() == 'null' else val

        return {
            "device_id":        clean(match.group(1)),
            "device_timestamp": clean(match.group(2)),
            "RS485_1":          clean(match.group(3)),
            "RS485_2":          clean(match.group(4)),
            "RS232_1":          clean(match.group(5)),
            "RS232_2":          clean(match.group(6)),
        }
    except Exception as e:
        print(f"❌ Parse error: {e}")
        return None


topics = [
    "gateway_Wifi_2/telemetry/CyclicData",
    "gateway_LTE_1/telemetry/CyclicData"
]

device_last_seen = {}
DEVICE_OFFLINE_SECONDS = 15


def on_connect(client, userdata, flags, rc):
    if rc == 0:
        print("✅ Connected to MQTT Broker!")
        for topic in topics:
            client.subscribe(topic, qos=1)
    else:
        print(f"❌ Connection failed with code {rc}")


def on_message(client, userdata, msg):
    try:
        raw = msg.payload.decode().strip()
        print(f"📡 MQTT → {msg.topic} : {raw}")

        data = parse_gateway_payload(raw)
        print("this is data --> ", data)

        device_id = msg.topic.split("/")[0]
        device_last_seen[device_id] = datetime.now()
        socketio.emit("device_status", {"device_id": device_id, "status": "online"})

        conn = connect_db()
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO gateway_data (
                device_id, rs485_1, rs485_2, rs232_1, rs232_2, device_timestamp
            ) VALUES (?, ?, ?, ?, ?, ?)
        """, (
            device_id, data["RS485_1"], data["RS485_2"],
            data["RS232_1"], data["RS232_2"], data["device_timestamp"],
        ))
        conn.commit()
        cursor.close()
        conn.close()
        print("✅ Data stored in DB")

        socketio.emit("sensor_update", {
            "device_id":        device_id,
            "device_timestamp": data["device_timestamp"],
            "RS485_1":          data["RS485_1"],
            "RS485_2":          data["RS485_2"],
            "RS232_1":          data["RS232_1"],
            "RS232_2":          data["RS232_2"],
        })
        print("📡 SOCKET → sensor_update sent")

        # ── Forward raw string to configured URL ──
        fwd_url = configured_url.get("url", "")
        if fwd_url:
            try:
                resp = requests.post(fwd_url, data=raw, headers={"Content-Type": "text/plain"}, timeout=5)
                print(f"✅ Forwarded to {fwd_url} → status {resp.status_code}")
            except Exception as fwd_err:
                print(f"⚠️ Forward failed: {fwd_err}")

    except Exception as e:
        print("❌ Error processing MQTT message:", e)


def connect_mqtt():
    client = mqtt.Client()
    client.username_pw_set(USERNAME, PASSWORD)
    client.on_connect = on_connect
    client.on_message = on_message
    try:
        print(f"Connecting to MQTT broker {BROKER}:{PORT} ...")
        client.connect(BROKER, PORT)
        client.loop_start()
    except Exception as e:
        print(f"❌ Exception occurred during connection: {e}")
    return client


@app.route("/", methods=["GET"])
def home():
    if "user_id" not in session:
        return redirect(url_for("login_page"))
    return redirect(url_for("devices_page"))


@app.route("/login", methods=["GET", "POST"])
def login_page():
    if request.method == "POST":
        data = request.get_json()
        email = data.get("email", "").strip()
        password = data.get("password", "")

        conn = connect_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM users WHERE email = ?", (email,))
        user = cursor.fetchone()
        cursor.close()
        conn.close()

        if user and check_password_hash(user["password"], password):
            session["user_id"] = user["id"]
            session["user_name"] = user["name"]
            return jsonify({"status": "success"})
        return jsonify({"status": "error", "message": "Invalid email or password"})

    if "user_id" in session:
        return redirect(url_for("devices_page"))
    return render_template("login.html")


@app.route("/register", methods=["POST"])
def register():
    data = request.get_json()
    name = data.get("name", "").strip()
    email = data.get("email", "").strip()
    password = data.get("password", "")
    rtu_no = data.get("rtu_no", "")
    temperature = data.get("temperature", 0)
    voltage = data.get("voltage", 0)
    current = data.get("current", 0)
    power = data.get("power", 0)

    if not name or not email or not password:
        return jsonify({"status": "error", "message": "All fields are required"})

    hashed_password = generate_password_hash(password)

    conn = connect_db()
    cursor = conn.cursor()
    try:
        cursor.execute("""
            INSERT INTO users (name, email, password, rtu_no, temperature, voltage, current, power)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (name, email, hashed_password, rtu_no, temperature, voltage, current, power))
        conn.commit()
        return jsonify({"status": "success"})
    except sqlite3.IntegrityError:
        return jsonify({"status": "error", "message": "Email already registered"})
    finally:
        cursor.close()
        conn.close()


@app.route("/logout", methods=["GET"])
def logout():
    session.clear()
    return redirect(url_for("login_page"))


@app.route("/devices", methods=["GET"])
def devices_page():
    if "user_id" not in session:
        return redirect(url_for("login_page"))
    return render_template("devices.html")


@app.route('/profile')
def profile():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    return render_template('profile.html')


@app.route("/device_status", methods=["GET"])
def device_status():
    if "user_id" not in session:
        return jsonify({}), 401
    now = datetime.now()
    status = {}
    for t in topics:
        device_id = t.split("/")[0]
        last_seen = device_last_seen.get(device_id)
        if last_seen and (now - last_seen).total_seconds() <= DEVICE_OFFLINE_SECONDS:
            status[device_id] = "online"
        else:
            status[device_id] = "offline"
    return jsonify(status)


@app.route("/dashboard", methods=["GET"])
def dashboard():
    if "user_id" not in session:
        return redirect(url_for("login_page"))
    device = request.args.get("device")
    if device:
        session["selected_device"] = device
    return render_template("index.html", device=session.get("selected_device", ""))


@app.route("/save_config", methods=["POST"])
def save_config():
    if "user_id" not in session:
        return jsonify({}), 401
    data = request.get_json()
    company = data.get("company", "").strip()
    url = data.get("url", "").strip()
    configured_url["url"] = url

    conn = connect_db()
    cursor = conn.cursor()
    # Keep only the latest config (upsert pattern)
    cursor.execute("DELETE FROM configured_urls")
    cursor.execute(
        "INSERT INTO configured_urls (company, url) VALUES (?, ?)",
        (company, url)
    )
    conn.commit()
    cursor.close()
    conn.close()

    print(f"✅ Configured URL saved to DB: {company} → {url}")
    return jsonify({"status": "success"})


@app.route("/graph_data", methods=["GET"])
def graph_data():
    start_date = request.args.get("start_date")
    end_date   = request.args.get("end_date")

    if not start_date or not end_date:
        start_date = end_date = datetime.now().strftime("%Y-%m-%d")

    start_str = start_date + " 00:00:00"
    end_str   = end_date   + " 23:59:59"

    device = request.args.get("device", "")

    conn = connect_db()
    cursor = conn.cursor()
    if device:
        cursor.execute("""
            SELECT device_timestamp, rs485_1, rs485_2, rs232_1, rs232_2
            FROM gateway_data
            WHERE device_timestamp >= ? AND device_timestamp <= ? AND device_id = ?
            ORDER BY device_timestamp ASC
        """, (start_str, end_str, device))
    else:
        cursor.execute("""
            SELECT device_timestamp, rs485_1, rs485_2, rs232_1, rs232_2
            FROM gateway_data
            WHERE device_timestamp >= ? AND device_timestamp <= ?
            ORDER BY device_timestamp ASC
        """, (start_str, end_str))

    rows = cursor.fetchall()
    cursor.close()
    conn.close()

    labels, rs485_1, rs485_2, rs232_1, rs232_2 = [], [], [], [], []
    for row in rows:
        labels.append(row["device_timestamp"])
        rs485_1.append(row["rs485_1"])
        rs485_2.append(row["rs485_2"])
        rs232_1.append(row["rs232_1"])
        rs232_2.append(row["rs232_2"])

    return jsonify({
        "labels":  labels,
        "RS485_1": rs485_1,
        "RS485_2": rs485_2,
        "RS232_1": rs232_1,
        "RS232_2": rs232_2,
    })


@app.route("/download", methods=["GET"])
def download():
    start = request.args.get("start")
    end   = request.args.get("end")

    conn = connect_db()
    cursor = conn.cursor()
    if start and end:
        cursor.execute("""
            SELECT device_id, rs485_1, rs485_2, rs232_1, rs232_2, device_timestamp
            FROM gateway_data
            WHERE device_timestamp >= ? AND device_timestamp <= ?
            ORDER BY device_timestamp ASC
        """, (start + " 00:00:00", end + " 23:59:59"))
    else:
        cursor.execute("""
            SELECT device_id, rs485_1, rs485_2, rs232_1, rs232_2, device_timestamp
            FROM gateway_data ORDER BY device_timestamp ASC
        """)
    rows = cursor.fetchall()
    cursor.close()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sensor Data"
    header_fill = PatternFill("solid", fgColor="1A1A2E")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    headers = ["Device ID", "RS485_1", "RS485_2", "RS232_1", "RS232_2", "Timestamp"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, 2):
        for col_idx, val in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"sensor_data_{start or 'all'}_to_{end or 'all'}.xlsx"
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.route("/latest", methods=["GET"])
def latest():
    device = request.args.get("device", "")

    conn = connect_db()
    cursor = conn.cursor()
    if device:
        cursor.execute("""
            SELECT device_id, rs485_1, rs485_2, rs232_1, rs232_2, device_timestamp
            FROM gateway_data
            WHERE device_id = ?
            ORDER BY id DESC
            LIMIT 1
        """, (device,))
    else:
        cursor.execute("""
            SELECT device_id, rs485_1, rs485_2, rs232_1, rs232_2, device_timestamp
            FROM gateway_data
            ORDER BY id DESC
            LIMIT 1
        """)
    row = cursor.fetchone()
    cursor.close()
    conn.close()

    if row is None:
        return jsonify({}), 204

    return jsonify({
        "device_id":        row["device_id"],
        "RS485_1":          row["rs485_1"],
        "RS485_2":          row["rs485_2"],
        "RS232_1":          row["rs232_1"],
        "RS232_2":          row["rs232_2"],
        "device_timestamp": row["device_timestamp"],
    })


if __name__ == "__main__":
    create_tables()
    load_configured_url()   # ← loads saved URL from DB into memory
    mqtt_client = connect_mqtt()
    socketio.run(app, host="0.0.0.0", port=5001, debug=True)